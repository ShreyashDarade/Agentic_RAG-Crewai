"""Review round 1 (parsers and chunker): document order, merged cells, size bounds, linear time."""

from __future__ import annotations

import io
import time
import zipfile

import pytest

pytest.importorskip("docx")

import docx

from agentic_rag.adapters.chunker_recursive import RecursiveChunker
from agentic_rag.adapters.parser_docx import DocxParser
from agentic_rag.errors import LimitExceeded
from agentic_rag.ports import ParsedDocument


def _docx(build) -> bytes:  # type: ignore[no-untyped-def]
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_table_stays_between_the_paragraphs_around_it() -> None:
    def build(d: docx.document.Document) -> None:
        d.add_paragraph("Section 1: Prices")
        table = d.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "widget"
        table.cell(0, 1).text = "9 EUR"
        d.add_paragraph("Section 2: Warranty")

    text = DocxParser().parse(_docx(build), name="a.docx").text
    assert text.index("Section 1") < text.index("widget | 9 EUR") < text.index("Section 2")


def test_a_merged_cell_is_read_once() -> None:
    def build(d: docx.document.Document) -> None:
        table = d.add_table(rows=2, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 2))
        merged.text = "Quarterly totals"
        for i in range(3):
            table.cell(1, i).text = f"q{i}"

    text = DocxParser().parse(_docx(build), name="a.docx").text
    assert text.count("Quarterly totals") == 1
    assert "q0 | q1 | q2" in text


def test_a_part_that_declares_too_much_xml_is_refused_before_parsing() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "<w:document>" + "a" * (3 * 1024 * 1024) + "</w:document>")
    with pytest.raises(LimitExceeded):
        DocxParser(max_member_bytes=1024 * 1024).parse(buffer.getvalue(), name="bomb.docx")


def test_the_defaults_bound_a_single_part_far_below_the_old_100_mib_total() -> None:
    parser = DocxParser()
    assert parser._max_member <= 8 * 1024 * 1024
    assert parser._max_bytes <= 32 * 1024 * 1024


def test_chunking_one_enormous_line_scales_linearly() -> None:
    """Quadratic behaviour shows as time growing with the square of the input; a wall-clock bound would depend on the machine."""
    chunker = RecursiveChunker(max_chars=1500, overlap_chars=200)

    def best_of_two(megabytes: int) -> float:
        text = "word " * (
            megabytes * 200_000
        )  # no paragraph or sentence breaks: everything goes through the hard split
        best = float("inf")
        for _ in range(2):
            started = time.perf_counter()
            chunks = chunker.chunk(ParsedDocument(text=text), document_id="d", document_name="n.txt")
            best = min(best, time.perf_counter() - started)
        assert max(len(c.text) for c in chunks) <= 1500
        return best

    small, large = best_of_two(2), best_of_two(8)  # four times the input
    ratio = large / small
    assert ratio < 12, f"4x the text took {ratio:.1f}x as long (about 4 is linear, about 16 or more is quadratic)"


def test_hard_split_loses_no_words_and_respects_the_bound() -> None:
    chunker = RecursiveChunker(max_chars=100, overlap_chars=10)
    words = [f"w{i}" for i in range(3000)]
    chunks = chunker.chunk(ParsedDocument(text=" ".join(words)), document_id="d", document_name="n.txt")
    covered = set(" ".join(c.text for c in chunks).split())
    assert set(words) <= covered
    assert all(len(c.text) <= 100 for c in chunks)
    unbroken = "x" * 1000
    pieces = chunker.chunk(ParsedDocument(text=unbroken), document_id="d", document_name="n.txt")
    assert "".join(c.text for c in pieces).count("x") >= 1000
    assert all(len(c.text) <= 100 for c in pieces)


def test_the_chunker_version_changes_with_its_settings() -> None:
    assert RecursiveChunker(max_chars=500).version != RecursiveChunker(max_chars=600).version
    assert RecursiveChunker(overlap_chars=10).version != RecursiveChunker(overlap_chars=20).version
    assert RecursiveChunker().version == RecursiveChunker().version


@pytest.mark.parametrize("unit", ["x" * 100, "ab " * 40, "same line\n" * 30])
def test_periodic_text_loses_nothing_even_when_the_last_chunk_equals_the_one_before(unit: str) -> None:
    chunker = RecursiveChunker(max_chars=100, overlap_chars=10)
    text = unit * 12
    chunks = chunker.chunk(ParsedDocument(text=text), document_id="d", document_name="n.txt")
    assert sum(len(c.text.replace("\n", "").replace(" ", "")) for c in chunks) >= len(
        text.replace("\n", "").replace(" ", "")
    )
